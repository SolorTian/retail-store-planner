"""Read-only personal oil summary. Department aggregates are never used."""
import io,re,urllib.parse
from html import unescape
import openpyxl
from datetime import datetime
from test_ttl_export import company_session

def download_fuel(account,password,year,month):
 request=company_session(account,password)
 url,body=request('/ubmsys_report_oil_expense/custom_lists?'+urllib.parse.urlencode({'task_str_date':f'{year}-{month:02d}'}))
 page=body.decode('utf-8')
 if '/login' in url:raise ValueError('Company login failed')
 staff=account.split('@')[0]
 matched=[]
 for row in re.findall(r'<tr\b[^>]*>.*?</tr>',page,re.S):
  text=re.sub('<[^>]+>',' ',unescape(row))
  if not re.search(r'-'+re.escape(staff)+r'(?=\s|$)',text):continue
  matched.extend(unescape(href) for href in re.findall(r'href=["\']([^"\']+)',row) if urllib.parse.urlsplit(unescape(href)).path=='/ubmsys_report_oil_expense/custom_form')
 if len(set(matched))!=1:raise ValueError('Personal mileage row unavailable')
 path=matched[0];split=urllib.parse.urlsplit(path);query=urllib.parse.parse_qs(split.query)
 if not query.get('user_id') or query.get('user_id')==['']:raise ValueError('Personal mileage row unavailable')
 _,detail=request(path)
 if b'/ubmsys_report_oil_expense/custom_form/export?' not in detail:raise ValueError('Mileage export unavailable')
 _,data=request('/ubmsys_report_oil_expense/custom_form/export?'+urllib.parse.urlencode(query,doseq=True))
 if not data.startswith(b'PK'):raise ValueError('Mileage export not XLSX')
 return data

def inspect_fuel(data):
 w=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True)
 try:return [[str(x) if x is not None else None for x in row] for row in list(w.active.iter_rows(values_only=True))[:12]]
 finally:w.close()

def parse_fuel(data,account,year,month):
 w=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True)
 try:
  rows=list(w.active.iter_rows(values_only=True))
  if len(rows)<6 or rows[3][0]!='日期' or rows[3][1]!='起迄地點' or rows[3][4]!='里程數(公里)':raise ValueError('Invalid mileage columns')
  owner=str(rows[1][1] or '').strip()
  if not owner.endswith('-'+account.split('@')[0]):raise ValueError('Personal mileage owner mismatch')
  items=[];seen=set()
  for row in rows[5:]:
   if not row[0]:continue
   date=row[0].strftime('%Y-%m-%d') if isinstance(row[0],datetime) else str(row[0]).strip()
   if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',date):continue
   if date[:7]!=f'{year}-{month:02d}' or date in seen:raise ValueError('Invalid mileage dates')
   seen.add(date);km=float(row[4])
   if not 0<=km<=10000:raise ValueError('Invalid mileage value')
   items.append({'date':date,'route':str(row[1] or '').strip()[:12000],'km':round(km,3)})
  if not items:raise ValueError('Empty personal mileage report')
  return {'version':1,'period':f'{year}-{month:02d}','owner':owner,'department':str(rows[2][1] or ''),'items':sorted(items,key=lambda x:x['date'])}
 finally:w.close()
