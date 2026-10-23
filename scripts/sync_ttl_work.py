"""Read-only monthly work export; data are encrypted by the sync entry point."""
import io
import math
import time
from concurrent.futures import ThreadPoolExecutor
import re
import urllib.parse
from datetime import datetime, timedelta
from html import unescape
import openpyxl
from test_ttl_export import company_session

FIELDS = {
 'id':['工作ID','工作 ID'], 'code':['店家代碼','客戶代碼'], 'name':['店家名稱','客戶名稱'],
 'owner':['業務員','執行人員'], 'channel':['店家型態','通路別'], 'grade':['客戶階級','階級'],
 'plannedAt':['工作可執行期間'], 'createdAt':['建立時間'], 'checkInAt':['打卡時間','到站時間'],
 'gpsCheckInAt':['到點GPS時間'],
 'completedAt':['工作完成日期'], 'checkOutAt':['離站時間','退卡時間'],
 'status':['任務狀態'], 'readStatus':['APP讀取狀態'], 'subject':['工作主題'], 'note':['備註'],
 'workflow':['工作流程','工作流程名稱','執行工作','問卷名稱','工作內容'],
 'signalValue':['是否衛星訊號異常','衛星訊號異常'], 'signalAbnormal':['是否衛星訊號異常','衛星訊號異常'], 'deviceFlag':['是否刷機'],
 'offsetAbnormal':['是否打卡偏移','是否偏移','打卡偏移'],
 'checkInLat':['打卡緯度','到站緯度'], 'checkInLng':['打卡經度','到站經度'],
 'checkOutLat':['離站緯度','退卡緯度'], 'checkOutLng':['離站經度','退卡經度'],
 'toleranceMeters':['容許範圍(公尺)','容許範圍（公尺）','容許距離(公尺)']
}
FLAGS={'signalAbnormal','deviceFlag','offsetAbnormal'}
NUMBERS={'checkInLat','checkInLng','checkOutLat','checkOutLng','toleranceMeters'}
def flag(value):
 text=str(value if value is not None else '').strip()
 if text in ('是','異常','1','True','true','Y','YES'): return True
 if text in ('否','正常','0','False','false','N','NO'): return False
 return None

def parse_work(data):
 workbook=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True)
 try:
  rows=iter(workbook.active.iter_rows(values_only=True)); header=None
  for _ in range(10):
   row=next(rows,None)
   if row is None: break
   if any(str(x or '').strip() in FIELDS['id'] for x in row):header=[str(x or '').strip() for x in row];break
  if header is None:raise ValueError('Invalid work header')
  ix={k:next((header.index(n) for n in names if n in header),None) for k,names in FIELDS.items()}
  if any(ix[k] is None for k in ['id','code','checkInAt','completedAt','owner']):raise ValueError('Missing work columns')
  items={}
  for row in rows:
   if len(row)<=ix['id'] or not row[ix['id']]:continue
   item={}
   for k,i in ix.items():
    value=row[i] if i is not None and i<len(row) else None
    if k=='signalAbnormal':item[k]=(str(value if value is not None else '').strip()!='否') if i is not None else None
    elif k=='signalValue':item[k]=str(value if value is not None else '').strip()[:300] if i is not None else None
    elif k in FLAGS:item[k]=flag(value)
    elif k in NUMBERS:
     try:item[k]=float(value) if value not in (None,'') else None
     except (ValueError,TypeError):item[k]=None
     if isinstance(value,bool) or item[k] is not None and (
      not math.isfinite(item[k]) or k.endswith('Lat') and abs(item[k])>90
      or k.endswith('Lng') and abs(item[k])>180):item[k]=None
    else:item[k]=value.strftime('%Y-%m-%d %H:%M:%S') if isinstance(value,datetime) else str(value or '').strip()[:2000 if k=='note' else 300]
   if not re.fullmatch(r'[A-Za-z0-9_-]{1,100}',item['id']):raise ValueError('Invalid work ID')
   items[item['id']]=item
   if len(items)>50000:raise ValueError('Work report too large')
  return {'policyVersion':2,'items':list(items.values()),'columns':[k for k,i in ix.items() if i is not None], 'timeBasis':'checkOutAt' if ix['checkOutAt'] is not None else 'completedAt'}
 finally:workbook.close()

def work_range(year,month):
 anchor=datetime(year,month,1)
 return (anchor-timedelta(days=1)).replace(day=1), (anchor.replace(day=28)+timedelta(days=4)).replace(day=1)-timedelta(days=1)

def parse_checkin_detail(page, owner):
 """Read observed task-message event markup; never substitute store coordinates."""
 events=[]
 for block in re.findall(r'<li\b[^>]*>(.*?)</li>',page,re.S|re.I):
  title=re.search(r'<div\b[^>]*class=["\'][^"\']*\btitle\b[^"\']*["\'][^>]*>(.*?)</div>',block,re.S|re.I)
  content=re.search(r'<div\b[^>]*class=["\'][^"\']*\bcontent\b[^"\']*["\'][^>]*>(.*?)</div>',block,re.S|re.I)
  if not title or not content:continue
  text=' '.join(unescape(re.sub(r'<[^>]+>',' ',title[1])).split())
  event=re.match(re.escape(owner)+r'\s+(離站打卡|到站打卡|打卡)\s',text)
  stamp=re.search(r'\(\s*(\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})\s*\)',text)
  if not event or not stamp:continue
  datetime.strptime(stamp[1],'%Y-%m-%d %H:%M:%S')
  body=unescape(re.sub(r'<[^>]+>',' ',content[1]))
  coordinate=re.search(r'經緯度\s*[:：]\s*([-+\d.]+)\s*[,，]\s*([-+\d.]+)',body)
  point=None
  if coordinate:
   try:
    lat,lng=map(float,coordinate.groups())
    if math.isfinite(lat) and math.isfinite(lng) and abs(lat)<=90 and abs(lng)<=180 and (lat,lng)!=(0,0):point=(lat,lng)
   except ValueError:pass
  events.append((stamp[1],event[1]=='離站打卡',point))
 arrivals=sorted((e for e in events if not e[1]),key=lambda e:e[0])
 if not arrivals:return {}
 first=arrivals[0]
 # Ambiguous duplicate events at the same instant cannot establish one GPS point.
 def unique_point(matches):
  points={e[2] for e in matches}
  return next(iter(points)) if len(points)==1 else None
 first_point=unique_point(e for e in arrivals if e[0]==first[0])
 result={'checkInAt':first[0],'checkInLat':first_point[0] if first_point else None,'checkInLng':first_point[1] if first_point else None}
 departures=sorted((e for e in events if e[1] and e[0]>=first[0]),key=lambda e:e[0])
 if departures:
  last=departures[-1];point=unique_point(e for e in departures if e[0]==last[0])
  result.update(checkOutAt=last[0],checkOutLat=point[0] if point else None,checkOutLng=point[1] if point else None)
 return result

def enrich_work_export(data, request, include_all=True, deadline_seconds=480):
 snapshot=parse_work(data)
 temporary={r['id']:r for r in snapshot['items'] if (include_all or r['name'].strip()=='臨時點') and (r['checkInAt'] or r['completedAt'])}
 if not temporary:return data
 deadline=time.monotonic()+deadline_seconds
 def fetch_detail(pair):
  task_id,row=pair
  if time.monotonic()>=deadline:return task_id,{}
  try:
   url,body=request('/ubmsys_task_message/lists/'+urllib.parse.quote(task_id,safe=''))
   page=body.decode('utf-8')
   if '/login' in url or 'id="user_pwd"' in page:raise ValueError('Company login failed')
   if '工作討論' not in page:return task_id,{}
   return task_id,parse_checkin_detail(page,row['owner'])
  except ValueError as error:
   if str(error)=='Company login failed':raise
   return task_id,{}
  except Exception:
   # Keep the report usable; an unavailable detail remains unknown, never normal.
   return task_id,{}
 with ThreadPoolExecutor(max_workers=4) as pool:
  details=dict(pool.map(fetch_detail,temporary.items()))
 for detail in details.values():
  if detail.get('checkInAt'):detail['gpsCheckInAt']=detail['checkInAt']
 book=openpyxl.load_workbook(io.BytesIO(data))
 try:
  sheet=book.active
  header_row=next(row for row in sheet.iter_rows(max_row=10) if any(str(c.value or '').strip() in FIELDS['id'] for c in row))
  header=[str(c.value or '').strip() for c in header_row];line=header_row[0].row
  id_col=next(header.index(n)+1 for n in FIELDS['id'] if n in header)
  keys=['checkInAt','gpsCheckInAt','checkInLat','checkInLng','checkOutAt','checkOutLat','checkOutLng'];indexes={}
  for key in keys:
   col=next((header.index(n)+1 for n in FIELDS[key] if n in header),None)
   if col is None:header.append(FIELDS[key][0]);col=len(header);sheet.cell(line,col,FIELDS[key][0])
   indexes[key]=col
  for row in sheet.iter_rows(min_row=line+1):
   task_id=str(row[id_col-1].value or '').strip()
   if task_id not in temporary:continue
   for key,col in indexes.items():
    if key=='checkInAt' and temporary[task_id]['name'].strip()!='臨時點':continue
    if key in details[task_id] or key not in ('checkInAt','checkOutAt'):sheet.cell(row[0].row,col).value=details[task_id].get(key)
  stream=io.BytesIO();book.save(stream);return stream.getvalue()
 finally:book.close()

def enrich_temporary_export(data,request):
 # Compatibility entry point for the original temporary-only tests.
 return enrich_work_export(data,request,include_all=False)

def scope_work_export(data,owner_staff):
 if not re.fullmatch(r'[0-9]{4,12}',owner_staff):raise ValueError('Invalid owner staff')
 book=openpyxl.load_workbook(io.BytesIO(data))
 try:
  sheet=book.active
  header=next(row for row in sheet.iter_rows(max_row=10) if any(str(c.value or '').strip() in FIELDS['id'] for c in row))
  columns=[str(c.value or '').strip() for c in header]
  owner=next(columns.index(name) for name in FIELDS['owner'] if name in columns)
  # Filter before fetching detail GPS, so another visitor cannot consume this
  # account's detail-query budget or appear in its eventual Google-linked data.
  for index in range(sheet.max_row,header[0].row,-1):
   if not re.search(r'(?:^|[-－])'+re.escape(owner_staff)+r'$',str(sheet.cell(index,owner+1).value or '').strip()):sheet.delete_rows(index)
  stream=io.BytesIO();book.save(stream);return stream.getvalue()
 finally:book.close()

def download_work(account,password,year,month,owner_staff=None):
 start,end=work_range(year,month)
 dates=start.strftime('%Y-%m-%d')+' ~ '+end.strftime('%Y-%m-%d')
 request=company_session(account,password)
 path='/ubmsys_task/lists?'+urllib.parse.urlencode({'period':'0','interval_type':'0','task_daterange':dates})
 url,body=request(path);page=body.decode('utf-8')
 if '/login' in url or 'id="user_pwd"' in page:raise ValueError('Company login failed')
 links=[urllib.parse.urljoin(url,unescape(x)) for x in re.findall(r'href=["\']([^"\']+)["\']',page)]
 link=next((x for x in links if urllib.parse.urlsplit(x).path=='/ubmsys_task/lists' and urllib.parse.parse_qs(urllib.parse.urlsplit(x).query).get('taskExport')==['3']),None)
 if not link:raise ValueError('Work export unavailable')
 split=urllib.parse.urlsplit(link);query=urllib.parse.parse_qs(split.query,keep_blank_values=True)
 query['task_daterange']=[dates];query['period']=['0'];query['interval_type']=['0']
 # The observed dynamic export includes check-in evidence and stable work IDs.
 query['taskExport']=['1']
 query['title']=['batch_delete,period_work_date,department,salesman,accompany_supervisor,read_status,executive_task,task_proccess,channel,channel_store_code,channel_store,customer_name,channel_store_address,customer_lvl,created_date,status,create_user,checkin_time,mobile_fake_gps,mobile_root,task_finish_date,task_note,task_id,management,']
 _,data=request(urllib.parse.urlunsplit((split.scheme,split.netloc,split.path,urllib.parse.urlencode(query,doseq=True),'')))
 if not data.startswith(b'PK'):raise ValueError('Work export not XLSX')
 return enrich_work_export(scope_work_export(data,owner_staff) if owner_staff is not None else data,request)